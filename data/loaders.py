"""لایهٔ خواندن Excel. هیچ منطق دامنه‌ای اینجا نیست."""

from openpyxl import load_workbook

_FA_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٫", "0123456789.")


def _clean(value):
    if value is None:
        return ""
    return str(value).replace("\u200f", "").strip()


def _to_int(value):
    text = _clean(value).translate(_FA_DIGITS)
    return int(float(text)) if text else None


def read_rows(path, sheet=None):
    """هر سطر Excel را به dict با کلیدهای سطر هدر تبدیل می‌کند."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook[sheet] if sheet else workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    header = [_clean(cell) for cell in next(iterator)]
    rows = []
    for raw in iterator:
        if all(_clean(cell) == "" for cell in raw):
            continue
        row = {}
        for index, key in enumerate(header):
            if key:
                row[key] = raw[index] if index < len(raw) else None
        rows.append(row)
    workbook.close()
    return rows


def load_intro(path):
    """INTRO.xlsx → {q_key: {"text": ..., "options": [...]}} با حفظ ترتیب."""
    intro = {}
    for row in read_rows(path):
        key = _clean(row.get("q_key"))
        if not key:
            continue
        question = intro.setdefault(key, {"text": _clean(row.get("q_text")), "options": []})
        question["options"].append(
            {
                "text": _clean(row.get("opt_text")),
                "role": _clean(row.get("role")) or None,
                "gender": _clean(row.get("gender")) or None,
                "children": _clean(row.get("children")) or None,
                "product_key": _clean(row.get("product_key")) or None,
            }
        )
    return intro


def load_questions(path):
    """QUESTIONS.xlsx → لیست مرتب سؤالات با گزینه‌ها و کد هر گزینه."""
    buckets = {}
    for row in read_rows(path):
        index = _to_int(row.get("q_index"))
        if index is None:
            continue
        question = buckets.setdefault(index, {"index": index, "text": _clean(row.get("q_text")), "options": []})
        question["options"].append(
            {
                "label": _clean(row.get("opt_label")),
                "text": _clean(row.get("opt_text")),
                "code": _to_int(row.get("code")),
            }
        )
    return [buckets[key] for key in sorted(buckets)]


def load_result_texts(path):
    """RESULT_TEXTS.xlsx → {code: {"title", "nickname", "description"}}"""
    results = {}
    for row in read_rows(path):
        code = _to_int(row.get("code"))
        if code is None:
            continue
        results[code] = {
            "title": _clean(row.get("title")),
            "nickname": _clean(row.get("nickname")),
            "description": _clean(row.get("description")),
        }
    return results


def load_result_images(path):
    """RESULT_IMAGES.xlsx → {code: filename}؛ نام فایل هرگز تبدیل نمی‌شود."""
    images = {}
    for row in read_rows(path):
        code = _to_int(row.get("code"))
        if code is None:
            continue
        images[code] = _clean(row.get("image"))
    return images


def load_products(path):
    """PRODUCTS.xlsx → {product_key: {"name", "url"}}"""
    products = {}
    for row in read_rows(path):
        key = _clean(row.get("product_key"))
        if not key:
            continue
        products[key] = {"name": _clean(row.get("name")), "url": _clean(row.get("url"))}
    return products