#!/usr/bin/env python3
# tools/build_excels.py

"""
Generate data/*.xlsx from "ربات اجرایی.docx".

Run with:  python -m tools.build_excels
"""

import html
import re
import shutil
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook

# Add project root to sys.path for importing config
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from config import DOCX_PATH

# ----------------------------------------------------------------------
# Constants
# ----------------------------------------------------------------------
DATA_DIR = Path(__file__).resolve().parents[1] / "data"
DATA_DIR.mkdir(exist_ok=True)

QUESTIONS_FILE   = DATA_DIR / "QUESTIONS.xlsx"
INTRO_FILE       = DATA_DIR / "INTRO.xlsx"
RESULT_TEXTS_FILE = DATA_DIR / "RESULT_TEXTS.xlsx"
RESULT_IMAGES_FILE = DATA_DIR / "RESULT_IMAGES.xlsx"
PRODUCTS_FILE    = DATA_DIR / "PRODUCTS.xlsx"

# Mapping of each question's correct code order (5 options per question)
CODE_MAP = [
    [4, 2, 5, 1, 3],  # Q1
    [1, 5, 3, 2, 4],  # Q2
    [3, 1, 4, 5, 2],  # Q3
    [5, 2, 4, 1, 3],  # Q4
    [2, 4, 5, 1, 3],  # Q5
]

OPTION_LABELS = ["الف", "ب", "ج", "د", "ه"]

# Personality types: code -> (title, nickname, image_filename)
# Codes 1..5 must all be present (they are used in CODE_MAP).
# The original code had only 1,4,5; placeholders are added for 2 and 3.
TYPES = {
    1: ("والد پشتیبان و حمایتگر", "برج مراقبت", "borj_moraghebat.jpg"),
    2: ("نوع دوم (پشتیبان)", "نام مستعار ۲", "type2.jpg"),          # placeholder
    3: ("نوع سوم (متعادل)", "نام مستعار ۳", "type3.jpg"),          # placeholder
    4: ("آزادی و محبت افراطی و سروری بیش از اندازهٔ فرزند", "خودپرداز", "khodpardaz.jpg"),
    5: ("والد پرخاشگر", "پرتابگر دمپایی", "partabgar_dampayi.jpg"),
}

# Product links 
PRODUCTS = {
    "single": ("دوره آموزشی تشکیل خانواده", "https://www.keraamat.ir/FProduct/81/"),
    "mother": ("دوره آموزشی تثبیت و تحکیم خانواده", "https://www.keraamat.ir/FProduct/82/"),
    "father": ("دوره آموزشی بهبود روابط خانوادگی", "https://www.keraamat.ir/FProduct/113/"),
    "girl": ("دوره جامع تربیت دخترانه", "https://www.keraamat.ir/FProduct/35/"),
    "boy": ("دوره جامع تربیت پسرانه", "https://www.keraamat.ir/FProduct/34/"),
    "both": ("دوره آموزشی تربیت جنسی و مدیریت عواطف", "https://www.keraamat.ir/FProduct/115/"),
}

# Intro questions (marital status and child gender)
INTRO_ROWS = [
    ("marital", "وضعیت تاهل خودتونو انتخاب کنین.", "عزب می‌باشم (مجرد)", "", "", "", "single"),
    ("marital", "وضعیت تاهل خودتونو انتخاب کنین.", "مامان هستم", "mother", "female", "", "mother"),
    ("marital", "وضعیت تاهل خودتونو انتخاب کنین.", "بابا هستم", "father", "male", "", "father"),
    ("child_gender", "جنسیت فرزندتون چیه؟", "دختر قندعسل", "", "", "girl", "girl"),
    ("child_gender", "جنسیت فرزندتون چیه؟", "پسر کاکل‌به‌سر", "", "", "boy", "boy"),
    ("child_gender", "جنسیت فرزندتون چیه؟", "جنسمون جوره، دختر و پسر", "", "", "both", "both"),
]

# Regex patterns for parsing DOCX
QUESTION_HEAD = re.compile(r"^\s*(?:سؤال|سوال)?\s*([۰-۹0-9]+)\s*[).:\-–]")
OPTION_HEAD   = re.compile(r"^\s*(الف|ب|ج|د|ه|هـ)\s*[).:\-–]\s*(.+)$")
FA_DIGITS     = str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")

# ----------------------------------------------------------------------
# Helper functions
# ----------------------------------------------------------------------
def docx_paragraphs(path: Path) -> list[str]:
    """
    Extract all text paragraphs from a .docx file (word/document.xml)
    without external dependencies (only zipfile and regex).
    """
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")

    paragraphs = []
    for block in re.split(r"</w:p>", xml):
        # Extract text from all <w:t> tags
        text = "".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", block, flags=re.S))
        text = html.unescape(re.sub(r"<[^>]+>", "", text))
        text = text.replace("\u200f", "").strip()
        if text:
            paragraphs.append(text)
    return paragraphs


def extract_questions(paragraphs: list[str]) -> list[dict]:
    """
    Group paragraphs into questions with their five options.
    Returns a list of dicts: {'text': question_stem, 'options': [opt1, ..., opt5]}.
    """
    groups = []
    current = None

    for line in paragraphs:
        option = OPTION_HEAD.match(line)
        if option and current is not None:
            current["options"].append(option.group(2).strip())
            continue

        # A numbered line or a long line starts a new question block
        if QUESTION_HEAD.match(line) or len(line) > 40:
            if current and len(current["options"]) == 5:
                groups.append(current)
            # Remove the question number from the stem
            stem = QUESTION_HEAD.sub("", line).strip()
            current = {"text": stem, "options": []}

    if current and len(current["options"]) == 5:
        groups.append(current)

    # Drop any group that is actually an intro question (by matching options)
    intro_texts = {row[2] for row in INTRO_ROWS}
    return [g for g in groups if not set(g["options"]) & intro_texts]


def write_sheet(path: Path, header: list[str], rows: list[tuple]) -> None:
    """
    Write an Excel sheet with given header and rows.
    Creates a backup (.bak) of any existing file.
    """
    if path.exists():
        shutil.copyfile(path, path.with_suffix(path.suffix + ".bak"))

    wb = Workbook()
    ws = wb.active

    ws.append(header)
    for row in rows:
        ws.append(list(row))

    # Set column widths
    for col_idx in range(1, len(header) + 1):
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[col_letter].width = 40

    wb.save(path)
    print(f"✓ Written: {path}")


# ----------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------
def main() -> None:
    if not DOCX_PATH.exists():
        sys.exit(f"Document not found: {DOCX_PATH}")

    paragraphs = docx_paragraphs(DOCX_PATH)
    questions = extract_questions(paragraphs)

    if len(questions) < len(CODE_MAP):
        sys.exit(
            f"Only {len(questions)} five-option questions were extracted, "
            f"but {len(CODE_MAP)} are required. "
            "Check the document structure or fill the question texts manually."
        )

    # Use only the first N questions (N = len(CODE_MAP))
    questions = questions[:len(CODE_MAP)]

    # Build rows for QUESTIONS.xlsx
    question_rows = []
    for q_index, (q_data, codes) in enumerate(zip(questions, CODE_MAP), start=1):
        for label, opt_text, code in zip(OPTION_LABELS, q_data["options"], codes):
            question_rows.append((q_index, q_data["text"], label, opt_text, code))

    write_sheet(
        QUESTIONS_FILE,
        ["q_index", "q_text", "opt_label", "opt_text", "code"],
        question_rows
    )

    write_sheet(
        INTRO_FILE,
        ["q_key", "q_text", "opt_text", "role", "gender", "children", "product_key"],
        INTRO_ROWS
    )

    # Result texts sheet (description is a placeholder)
    result_text_rows = []
    for code, (title, nickname, _) in sorted(TYPES.items()):
        description = f"{title} — «{nickname}» (توضیحات کامل را اینجا وارد کنید)"
        result_text_rows.append((code, title, nickname, description))
    write_sheet(
        RESULT_TEXTS_FILE,
        ["code", "title", "nickname", "description"],
        result_text_rows
    )

    # Result images sheet
    write_sheet(
        RESULT_IMAGES_FILE,
        ["code", "image"],
        [(code, img) for code, (_, _, img) in sorted(TYPES.items())]
    )

    # Products sheet
    write_sheet(
        PRODUCTS_FILE,
        ["product_key", "name", "url"],
        [(key, name, url) for key, (name, url) in PRODUCTS.items()]
    )

    print("\nReminder: the description column in RESULT_TEXTS.xlsx is still a stub; "
          "paste the full text for each personality type there.")


if __name__ == "__main__":
    main()