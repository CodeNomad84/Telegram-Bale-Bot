"""Offline self-test: imports every module and exercises the key flows.

Run from the Quizbot folder:  python verify.py

- Uses a temporary SQLite database (the real keramat.db is never touched).
- The bot's network methods are replaced by recorders, so nothing is sent to Bale.
- Exit code 0 = all checks passed, 1 = at least one failure.
"""

import asyncio
import importlib
import os
import sys
import tempfile
import traceback
from pathlib import Path
from types import SimpleNamespace

sys.stdout.reconfigure(encoding="utf-8")
sys.dont_write_bytecode = True
BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))

# ---------- temp database / working dir (must happen before other imports) ----------
import config  # noqa: E402

TMP = Path(tempfile.mkdtemp(prefix="quizbot_verify_"))
config.DB_DIR = TMP
config.DB_PATH = TMP / "verify.db"
os.chdir(TMP)  # export_*.xlsx, graph_*.png and scheduler_state.txt land here

RESULTS = []


def check(name, condition, detail=""):
    RESULTS.append((name, bool(condition)))
    mark = "PASS" if condition else "FAIL"
    print(f"[{mark}] {name}" + (f"  ({detail})" if detail and not condition else ""))


# ---------- 1) import every module ----------
MODULES = [
    "config", "bot_instance", "handlers.states", "database.db", "data.loaders", "data.content",
    "keyboards.builders", "utils.helpers", "utils.decorators", "handlers.membership",
    "handlers.intro", "handlers.quiz", "handlers.result", "handlers.registration",
    "handlers.prize_info", "admin.stats", "admin.broadcast", "admin.export", "admin.graph",
    "admin.prize", "admin.logs", "admin.handlers", "scheduler", "main",
]
for module_name in MODULES:
    try:
        importlib.import_module(module_name)
        check(f"import {module_name}", True)
    except Exception as e:  # noqa: BLE001
        check(f"import {module_name}", False, f"{type(e).__name__}: {e}")

if not all(ok for _, ok in RESULTS):
    print("\nImport failures; aborting further checks.")
    sys.exit(1)

import main  # noqa: E402
import handlers.membership as membership  # noqa: E402
from admin import handlers as admin_handlers  # noqa: E402
from admin.export import export_users_excel, export_winners_excel  # noqa: E402
from admin.graph import generate_graph  # noqa: E402
from admin.logs import get_logs_text  # noqa: E402
from admin.prize import format_winners_list, lottery_top_inviters  # noqa: E402
from admin.stats import get_stats_text  # noqa: E402
from bot_instance import bot  # noqa: E402
from data import content  # noqa: E402
from database import db  # noqa: E402
from handlers.prize_info import handle_winner_info  # noqa: E402
from handlers.result import (  # noqa: E402
    build_description, build_products, compute_code, invite_link, result_image_path, send_result,
)
from handlers.states import (  # noqa: E402
    STATE_AWAIT_JOIN, STATE_AWAIT_PHONE, STATE_DONE, STATE_INTRO, STATE_QUIZ,
    build_display_name, can_answer, normalize_username,
)
from utils.helpers import safe_answer, to_persian_digits  # noqa: E402

# ---------- fake bot (records calls instead of hitting the API) ----------
SENT = []
MEMBER_STATUS = {"value": "member"}


async def fake_send_message(chat_id, text=None, reply_markup=None, **kwargs):
    SENT.append(("message", chat_id, text))


async def fake_send_photo(chat_id, photo=None, caption=None, reply_markup=None, **kwargs):
    SENT.append(("photo", chat_id, caption))


async def fake_send_document(chat_id, document=None, caption=None, **kwargs):
    SENT.append(("document", chat_id, caption))


async def fake_get_chat_member(chat_id, user_id):
    return SimpleNamespace(status=MEMBER_STATUS["value"])


bot.send_message = fake_send_message
bot.send_photo = fake_send_photo
bot.send_document = fake_send_document
bot.get_chat_member = fake_get_chat_member
membership.CHANNEL_ID = "@verify_channel"  # force the membership check on


class FakeCallback:
    def __init__(self, user_id, data):
        self.author = SimpleNamespace(id=user_id, first_name="Test", last_name="User", username=None)
        self.data = data
        self.answers = []

    async def answer(self, text, show_alert=False):
        self.answers.append((text, show_alert))


def fake_message(user_id, text):
    author = SimpleNamespace(id=user_id, first_name="Test", last_name="User", username=None)
    return SimpleNamespace(author=author, chat=SimpleNamespace(type="private"), text=text)


async def tap(user_id, data):
    cq = FakeCallback(user_id, data)
    await main.on_callback_query.callback(cq)
    return cq


async def say(user_id, text):
    await main.on_message.callback(fake_message(user_id, text))


def sent_to(user_id, kind=None):
    return [s for s in SENT if s[1] == user_id and (kind is None or s[0] == kind)]


# ---------- 2) pure helpers ----------
def test_pure_functions():
    check("compute_code([1,1,2,3,4]) == 1", compute_code([1, 1, 2, 3, 4]) == 1)
    check("compute_code([]) is None", compute_code([]) is None)
    check("to_persian_digits(123) == '۱۲۳'", to_persian_digits(123) == "۱۲۳")
    check("normalize_username('@test_user')", normalize_username("@test_user") == "@test_user")
    check("normalize_username('ab') is None", normalize_username("ab") is None)
    author = SimpleNamespace(first_name="علی", last_name="رضایی", username="ali")
    check("build_display_name(full name)", build_display_name(author) == "علی رضایی")
    check("build_display_name(username only)",
          build_display_name(SimpleNamespace(first_name=None, last_name=None, username="ali")) == "@ali")
    check("build_display_name(empty)", build_display_name(SimpleNamespace()) == "")
    check("parse_start_payload('/start 42')", main.parse_start_payload("/start 42") == 42)
    check("parse_start_payload('/start abc')", main.parse_start_payload("/start abc") is None)
    user = {"state": STATE_QUIZ, "step": 2}
    check("can_answer: matching", can_answer(user, STATE_QUIZ, 2))
    check("can_answer: wrong step", not can_answer(user, STATE_QUIZ, 1))
    check("can_answer: wrong state", not can_answer(user, STATE_INTRO, 2))
    check("can_answer: None user", not can_answer(None, STATE_QUIZ, 0))
    check("invite_link contains ?start=", "?start=77" in invite_link(77))


# ---------- 3) content files ----------
def test_content():
    check("INTRO has marital & child_gender", set(content.INTRO_ORDER) <= set(content.INTRO))
    check("QUESTIONS has 5 questions", len(content.QUESTIONS) == 5)
    check("every question has 5 options with codes 1..5",
          all(sorted(o["code"] for o in q["options"]) == [1, 2, 3, 4, 5] for q in content.QUESTIONS))
    check("RESULT_TEXTS covers codes 1..5", set(content.RESULT_TEXTS) >= {1, 2, 3, 4, 5})
    check("RESULT_IMAGES covers codes 1..5", set(content.RESULT_IMAGES) >= {1, 2, 3, 4, 5})
    missing = []
    for code, image in content.RESULT_IMAGES.items():
        name, ext = os.path.splitext(image)
        for suffix in ("female", "male"):
            if not (config.PICS_DIR / f"{name}_{suffix}{ext}").exists():
                missing.append(f"{name}_{suffix}{ext}")
    check("all gendered result images exist", not missing, ", ".join(missing))
    intro_keys = {o["product_key"] for q in content.INTRO.values() for o in q["options"] if o["product_key"]}
    check("every intro product_key exists in PRODUCTS", intro_keys <= set(content.PRODUCTS),
          str(intro_keys - set(content.PRODUCTS)))
    check("every marital option has a product_key",
          all(o["product_key"] for o in content.INTRO["marital"]["options"]))
    check("all product URLs start with https://",
          all(p["url"].startswith("https://") for p in content.PRODUCTS.values()))
    longest = max(len(build_description(code)) for code in content.RESULT_TEXTS)
    check("result captions fit in 1024 chars", longest <= 1024, f"longest={longest}")
    check("build_products formats known keys", "🔹" in build_products(["single"]))
    check("build_products ignores unknown keys", build_products(["nope"]) == "")


def test_result_image_path():
    female_single = {"role": "single", "gender": "female"}
    father = {"role": "father", "gender": "male"}
    path = result_image_path(1, female_single)
    check("single female -> *_female image", path is not None and path.name.endswith("_female.png"))
    path = result_image_path(1, father)
    check("father -> *_male image", path is not None and path.name.endswith("_male.png"))
    check("unknown code -> None (no NameError)", result_image_path(999, father) is None)


# ---------- 4) database ----------
async def test_database():
    await db.init_db()
    check("ensure_user creates -> True", await db.ensure_user(1001, "Inviter") is True)
    check("ensure_user existing -> False", await db.ensure_user(1001, "Inviter 2") is False)
    user = await db.get_user(1001)
    check("ensure_user updates display_name", user["display_name"] == "Inviter 2")
    check("get_user unknown -> None", await db.get_user(424242) is None)

    await db.update_user(1001, answers=[1, 2, 3], product_keys=["single", "girl"])
    user = await db.get_user(1001)
    check("answers JSON roundtrip", user["answers"] == [1, 2, 3])
    check("product_keys JSON roundtrip", user["product_keys"] == ["single", "girl"])
    await db.reset_quiz(1001)
    user = await db.get_user(1001)
    check("reset_quiz clears answers/step", user["answers"] == [] and user["step"] == 0)

    await db.ensure_user(1002, "Friend")
    check("register_invite: unknown inviter rejected", await db.register_invite(555555, 1002) is False)
    check("register_invite: self rejected", await db.register_invite(1002, 1002) is False)
    check("register_invite: valid -> True", await db.register_invite(1001, 1002) is True)
    check("register_invite: duplicate -> False", await db.register_invite(1001, 1002) is False)
    check("get_invite_count == 1", await db.get_invite_count(1001) == 1)

    await db.log_admin_action(1, "test", "details")
    logs = await db.get_admin_logs(10)
    check("admin logs stored", logs and logs[0]["action"] == "test")
    check("get_all_user_ids", set(await db.get_all_user_ids()) >= {1001, 1002})

    await db.register_lottery_winner(1002, "random_participants")
    winner = await db.get_lottery_winner(1002)
    check("register/get lottery winner", winner and winner["category"] == "random_participants")
    await db.save_winner_info(1002, "Name - Address")
    winner = await db.get_lottery_winner(1002)
    check("save_winner_info", winner["info_text"] == "Name - Address")
    rows = await db.get_all_lottery_winners()
    check("get_all_lottery_winners joins users", rows and rows[0]["display_name"] == "Friend")


# ---------- 5) full user flow ----------
async def test_user_flow():
    uid = 2001
    await say(uid, "/start 1001")
    user = await db.get_user(uid)
    check("/start -> await_join", user["state"] == STATE_AWAIT_JOIN)
    check("new user's inviter is pending", main.PENDING_INVITER.get(uid) == 1001)

    MEMBER_STATUS["value"] = "left"
    cq = await tap(uid, "check_join")
    check("check_join while not member -> alert", cq.answers and cq.answers[0][1] is True)
    MEMBER_STATUS["value"] = "member"
    cq = await tap(uid, "check_join")
    check("check_join answered exactly once", len(cq.answers) == 1)
    check("invite registered after join", await db.get_invite_count(1001) == 2)
    check("state -> await_phone", (await db.get_user(uid))["state"] == STATE_AWAIT_PHONE)

    await say(uid, "09120000000")
    user = await db.get_user(uid)
    check("phone saved, intro started", user["phone"] == "09120000000" and user["state"] == STATE_INTRO)

    cq = await tap(uid, "intro:0:3")  # آقای متاهل -> father
    user = await db.get_user(uid)
    check("parent goes to child-gender question", user["role"] == "father" and user["step"] == 1)
    check("intro answer answered exactly once", len(cq.answers) == 1)

    cq = await tap(uid, "intro:0:0")  # old message
    user = await db.get_user(uid)
    check("stale intro tap rejected", user["role"] == "father" and user["step"] == 1
          and cq.answers[0][0] == main.STALE_ANSWER_TEXT)

    MEMBER_STATUS["value"] = "left"
    cq = await tap(uid, "intro:1:0")
    check("non-member answer rejected", (await db.get_user(uid))["state"] == STATE_INTRO
          and cq.answers[0][1] is True)
    MEMBER_STATUS["value"] = "member"

    await tap(uid, "intro:1:0")  # girl
    user = await db.get_user(uid)
    check("quiz started after intro", user["state"] == STATE_QUIZ and user["step"] == 0)
    check("product keys collected", user["product_keys"] == ["father", "girl"])

    # double tap on the same question at the same time
    await asyncio.gather(tap(uid, "quiz:0:0"), tap(uid, "quiz:0:1"))
    user = await db.get_user(uid)
    check("double tap records only one answer", len(user["answers"]) == 1 and user["step"] == 1,
          f"answers={user['answers']}")

    for q in range(1, 5):
        await tap(uid, f"quiz:{q}:0")
    user = await db.get_user(uid)
    check("quiz finished -> done", user["state"] == STATE_DONE and user["result_code"] in range(1, 6))
    photos = sent_to(uid, "photo")
    check("result image sent with description caption",
          photos and photos[0][2].startswith("🎯"))

    cq = await tap(uid, "quiz:4:0")
    check("tap after done rejected", len((await db.get_user(uid))["answers"]) == 5)

    await say(uid, "/start")
    cq = await tap(uid, "quiz:0:0")  # old quiz button after restart (gate bypass bug)
    user = await db.get_user(uid)
    check("old quiz button after /start rejected", user["state"] == STATE_AWAIT_JOIN and user["answers"] == [])

    await say(uid, "/start 1002")
    check("existing user's invite link not counted", uid not in main.PENDING_INVITER)

    SENT.clear()
    await db.update_user(uid, state=STATE_DONE)
    await say(uid, "hello")
    fallbacks = [s for s in sent_to(uid, "message") if s[2] == config.FALLBACK_TEXT]
    check("free text gets exactly one fallback reply", len(fallbacks) == 1, f"got {len(fallbacks)}")

    # result without any image -> description as text
    saved = dict(content.RESULT_IMAGES)
    content.RESULT_IMAGES.clear()
    try:
        await db.update_user(uid, answers=[1, 1, 1, 1, 1])
        SENT.clear()
        await send_result(uid)
        texts = [s for s in sent_to(uid, "message") if s[2] and s[2].startswith("🎯")]
        check("no image -> description sent as text", len(texts) == 1)
    finally:
        content.RESULT_IMAGES.update(saved)


# ---------- 6) admin flow ----------
async def test_admin():
    admin_id = 900001
    config.ADMIN_IDS.add(admin_id)  # same set object used by main and decorators
    await db.ensure_user(admin_id, "Admin")

    cq = await tap(3001, "admin:stats")
    check("non-admin callback gets alert", cq.answers == [("⛔ شما دسترسی ادمین ندارید.", True)])

    for data in ("admin:main", "admin:stats", "admin:logs", "admin:prize", "admin:lottery"):
        cq = await tap(admin_id, data)
        check(f"{data} answered exactly once", len(cq.answers) == 1, str(cq.answers))

    # ranking order: create inviters with different counts
    for inviter, n in ((4001, 1), (4002, 3), (4003, 2)):
        await db.ensure_user(inviter, f"Inviter{inviter}")
        for k in range(n):
            invited = inviter * 10 + k
            await db.ensure_user(invited, "x")
            await db.register_invite(inviter, invited)
    top = await lottery_top_inviters(10)
    check("top inviters ordered by count", top[:3] == [4002, 1001, 4003] or top[0] == 4002, str(top))
    text = await format_winners_list([4002, 4003, 4001], "top_inviters")
    lines = [l for l in text.splitlines() if l[:1].isdigit()]
    check("winners list keeps ranking order",
          "Inviter4002" in lines[0] and "Inviter4003" in lines[1] and "Inviter4001" in lines[2])
    check("winners list shows invite counts", "3 دعوت" in lines[0] and "1 دعوت" in lines[2])

    # prize message without a prior lottery -> confirmation first
    admin_handlers.ADMIN_LOTTERY_RESULTS.clear()
    cq = await tap(admin_id, "admin:prize_cat:top_inviters")
    state = admin_handlers.ADMIN_STATE.get(admin_id, {})
    check("auto lottery asks for confirmation", state.get("action") == "prize_confirm_winners")
    winners = list(state.get("winners", []))
    await tap(admin_id, "admin:prize_confirm")
    check("confirm -> waiting for prize text",
          admin_handlers.ADMIN_STATE.get(admin_id, {}).get("action") == "prize_message_text")
    await say(admin_id, "🎉 You won!")
    check("prize sent and state cleared", admin_id not in admin_handlers.ADMIN_STATE)
    check("winners registered in lottery_winners", await db.get_lottery_winner(winners[0]) is not None)

    # lottery via count, then prize -> results cleared after sending
    await tap(admin_id, "admin:lottery:random_inviters")
    await say(admin_id, "2")
    check("lottery result stored",
          "random_inviters" in admin_handlers.ADMIN_LOTTERY_RESULTS.get(admin_id, {}))
    await tap(admin_id, "admin:prize_cat:random_inviters")
    await say(admin_id, "prize text")
    check("lottery result cleared after sending",
          "random_inviters" not in admin_handlers.ADMIN_LOTTERY_RESULTS.get(admin_id, {}))

    # commands escape a pending admin action
    await tap(admin_id, "admin:broadcast")
    await say(admin_id, "/start")
    check("/start clears pending admin action", admin_id not in admin_handlers.ADMIN_STATE)

    # winner info collection
    await db.register_lottery_winner(4002, "top_inviters")
    check("winner info captured", await handle_winner_info(4002, "My address") is True)
    check("winner info only once", await handle_winner_info(4002, "again") is False)

    # reports
    path = await export_users_excel()
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    check("export has 3 sheets", wb.sheetnames == ["Users", "Inviters", "Lottery_Winners"], str(wb.sheetnames))
    wb.close()
    path, total, submitted = await export_winners_excel()
    wb = load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    check("winners export: RTL sheet with all winners", ws.sheet_view.rightToLeft and len(rows) == total)
    check("winners export: submitted rows first", all(r[7] for r in rows[:submitted])
          and not any(r[7] for r in rows[submitted:]))
    check("winners export: info text present", any(r[7] == "My address" for r in rows))
    wb.close()
    SENT.clear()
    cq = await tap(admin_id, "admin:winners_export")
    docs = sent_to(admin_id, "document")
    check("admin:winners_export sends one document", len(cq.answers) == 1 and len(docs) == 1
          and "مشخصات برندگان" in docs[0][2])
    check("stats text", "کل کاربران" in await get_stats_text())
    check("logs text", "لاگ" in await get_logs_text())
    check("graph image generated", Path(await generate_graph()).exists())


# ---------- 7) misc ----------
async def test_misc():
    class Boom:
        async def answer(self, text, show_alert=False):
            raise RuntimeError("query is too old")
    try:
        await safe_answer(Boom(), "")
        check("safe_answer swallows API errors", True)
    except Exception:  # noqa: BLE001
        check("safe_answer swallows API errors", False)

    import scheduler
    first, second = scheduler.get_next_message(), scheduler.get_next_message()
    check("scheduler rotates messages", first == scheduler.MESSAGES[0] and second == scheduler.MESSAGES[1])


async def run_async_tests():
    for test in (test_database, test_user_flow, test_admin, test_misc):
        try:
            await test()
        except Exception as e:  # noqa: BLE001
            traceback.print_exc()
            check(f"{test.__name__} crashed", False, f"{type(e).__name__}: {e}")


if __name__ == "__main__":
    test_pure_functions()
    test_content()
    test_result_image_path()
    loop = bot.dispatcher.event_loop  # same loop the bot uses in production
    asyncio.set_event_loop(loop)
    loop.run_until_complete(run_async_tests())

    failed = [name for name, ok in RESULTS if not ok]
    print(f"\n{len(RESULTS) - len(failed)}/{len(RESULTS)} checks passed. Temp dir: {TMP}")
    if failed:
        print("Failed:\n  - " + "\n  - ".join(failed))
    sys.exit(1 if failed else 0)
