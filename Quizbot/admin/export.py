"""Excel exports: full users/inviters/winners dump, and a formatted winners sheet."""

import json
import aiosqlite
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

from config import DB_PATH
from admin.prize import CATEGORY_NAMES
from database.db import get_all_lottery_winners


async def export_users_excel():
    """تولید فایل اکسل با سه شیت: کاربران، دعوت‌کنندگان، برندگان."""
    file_path = f"export_users_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb = Workbook()

    # ---------- شیت ۱: کاربران ----------
    ws = wb.active
    ws.title = "Users"
    ws.append([
        "user_id", "username", "display_name", "phone", "role", "gender",
        "children", "answers", "result_code", "result_title", "state", "created_at"
    ])
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        cursor = await db.execute(
            "SELECT user_id, username, display_name, phone, role, gender, children, "
            "answers, result_code, result_title, state, created_at FROM users"
        )
        for row in await cursor.fetchall():
            answers = json.loads(row["answers"]) if row["answers"] else []
            ws.append([
                row["user_id"], row["username"], row["display_name"], row["phone"],
                row["role"], row["gender"], row["children"],
                ", ".join(str(a) for a in answers),
                row["result_code"], row["result_title"], row["state"], row["created_at"],
            ])

    # ---------- شیت ۲: دعوت‌کنندگان ----------
    ws2 = wb.create_sheet("Inviters")
    ws2.append(["inviter_id", "username", "display_name", "phone", "invite_count"])
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        cursor = await db.execute(
            """SELECT i.inviter_id, i.count, u.username, u.display_name, u.phone
               FROM invites i
               LEFT JOIN users u ON u.user_id = i.inviter_id
               ORDER BY i.count DESC"""
        )
        for row in await cursor.fetchall():
            ws2.append([
                row["inviter_id"], row["username"], row["display_name"],
                row["phone"], row["count"],
            ])

    # ---------- شیت ۳: برندگان قرعه‌کشی ----------
    ws3 = wb.create_sheet("Lottery_Winners")
    ws3.append([
        "user_id", "username", "display_name", "phone",
        "category", "prize_sent_at", "info_text", "info_collected_at"
    ])
    winners = await get_all_lottery_winners()
    for w in winners:
        ws3.append([
            w["user_id"], w["username"], w["display_name"], w["phone"],
            w["category"], w["prize_sent_at"],
            w["info_text"], w["info_collected_at"],
        ])

    wb.save(file_path)
    return file_path

# ---------- خروجی مشخصات برندگان (منوی جوایز) ----------
WINNER_HEADERS = [
    ("ردیف", 7),
    ("نام", 22),
    ("نام کاربری", 18),
    ("شماره تماس", 16),
    ("دسته جایزه", 26),
    ("ارسال پیام جایزه", 18),
    ("وضعیت", 16),
    ("مشخصات ارسالی برنده", 55),
    ("زمان ثبت مشخصات", 18),
    ("شناسه کاربر", 14),
]
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_PENDING_FILL = PatternFill("solid", fgColor="FFF2CC")
_BORDER = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def _fmt_time(value):
    """2026-10-13T14:05:09 → 2026-10-13 14:05"""
    return value.replace("T", " ")[:16] if value else ""


async def export_winners_excel():
    """فایل اکسل مرتب از جدول lottery_winners برای ادمین.

    خروجی: (مسیر فایل، تعداد کل، تعداد ثبت‌کرده) یا (None, 0, 0) اگر برنده‌ای نباشد.
    """
    winners = await get_all_lottery_winners()
    if not winners:
        return None, 0, 0
    # ابتدا کسانی که مشخصات فرستاده‌اند (جدیدترین بالا)، سپس منتظرها
    submitted = sorted((w for w in winners if w["info_text"]),
                       key=lambda w: w["info_collected_at"] or "", reverse=True)
    pending = [w for w in winners if not w["info_text"]]

    wb = Workbook()
    ws = wb.active
    ws.title = "مشخصات برندگان"
    ws.sheet_view.rightToLeft = True

    ws.append([title for title, _ in WINNER_HEADERS])
    for col, (_, width) in enumerate(WINNER_HEADERS, 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 28

    for i, w in enumerate(submitted + pending, 1):
        has_info = bool(w["info_text"])
        ws.append([
            i,
            w["display_name"] or "",
            w["username"] or "",
            w["phone"] or "",
            CATEGORY_NAMES.get(w["category"], w["category"]),
            _fmt_time(w["prize_sent_at"]),
            "✅ ثبت شده" if has_info else "⏳ در انتظار",
            w["info_text"] or "",
            _fmt_time(w["info_collected_at"]),
            w["user_id"],
        ])
        for cell in ws[ws.max_row]:
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="center" if cell.column != 8 else "right")
            if not has_info:
                cell.fill = _PENDING_FILL

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    file_path = f"export_winners_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(file_path)
    return file_path, len(winners), len(submitted)
